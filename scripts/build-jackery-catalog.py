"""Build the reviewed, source-backed Jackery catalog. Never writes source files or DBs.

Run with: python -X utf8 scripts/build-jackery-catalog.py
Requires openpyxl only for read-only OOXML workbook access.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import openpyxl

REPO = Path(__file__).resolve().parents[1]
DEFAULT_INDEX = Path(r"D:\Codex\JGPD\outputs\jackery_bom_20261027\bom_data.json")
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
VERSION_KEYS = {"未标注版本": "default", "原始版": "original", "降本版": "costdown", "A版": "a", "B版": "b"}
SOFTWARE = {"36-0001-000118", "36-0003-000009"}


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def save_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def paragraphs(element):
    # Runs within a paragraph concatenate; paragraphs and explicit line breaks survive.
    result = []
    for para in element.findall(".//a:p", NS):
        result.append("".join("\n" if x.tag == "{" + NS["a"] + "}br" else (x.text or "")
                              for x in para.iter() if x.tag in {"{" + NS["a"] + "}t", "{" + NS["a"] + "}br"}))
    return "\n".join(result).strip()


def extract_ppt(path):
    slides = []
    with ZipFile(path) as archive:
        names = sorted((x for x in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", x)),
                       key=lambda x: int(re.search(r"slide(\d+)", x).group(1)))
        for name in names:
            xml = ET.fromstring(archive.read(name))
            titles = [paragraphs(x) for x in xml.findall(".//p:sp", NS)]
            tables = []
            for table in xml.findall(".//a:tbl", NS):
                rows = []
                for row_no, row in enumerate(table.findall("a:tr", NS), 1):
                    rows.append({"row": row_no, "cells": [
                        {"text": paragraphs(cell), "merge": dict(cell.attrib)}
                        for cell in row.findall("a:tc", NS)]})
                tables.append(rows)
            slides.append({"slide": int(re.search(r"slide(\d+)", name).group(1)),
                           "titles": [x for x in titles if x], "tables": tables})
    return slides


def verify_bom_source(source, records):
    """Recheck exact source cells used by the catalog against the archived index."""
    path = Path(source["path"])
    assert sha256(path) == source["sha256"], f"Source hash changed: {path}"
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in source["sheets"]:
            ws = book[sheet["sheet"]]
            values = list(ws.iter_rows(values_only=True))
            headers = values[sheet["header_row"] - 1]
            for record in records:
                if record["来源工作表"] != ws.title:
                    continue
                row = dict(zip(headers, values[record["来源行号"] - 1]))
                pairs = [("物料编码", "子项物料编码"), ("物料名称", "物料名称"),
                         ("规格型号", "规格型号"), ("标准用量（原表）", "标准用量"),
                         ("型号（原表）", "型号"), ("物料生命周期", "物料生命周期"), ("发料方式", "发料方式")]
                if source["model"] == "E1000":
                    pairs = [("物料编码", "料号(Part No.)(必填)"), ("物料名称", "物料名称(Name)(必填)"),
                             ("规格型号", "规格型号(Specification)"), ("型号（原表）", "机型(Model)")]
                for indexed, original in pairs:
                    assert record[indexed] == row.get(original), (path.name, record["来源行号"], indexed)
    finally:
        book.close()


def build(index_path, source_root):
    index = json.loads(index_path.read_text(encoding="utf-8"))
    root = source_root or Path(index["source_root"])
    records = [dict(zip(index["headers"], row)) for row in index["rows"]]
    catalog = {"schemaVersion": 1, "batchKey": "jackery-eu-20261030", "sources": [],
               "models": [], "excluded": [], "reviewRows": []}
    extraction = []
    source_files = {}

    def register(path, kind):
        relative = path.relative_to(root).as_posix()
        if relative not in source_files:
            entry = {"file": path.name, "relativePath": relative, "sha256": sha256(path), "kind": kind}
            source_files[relative] = entry
            catalog["sources"].append(entry)
        return source_files[relative]

    def reference(entry, **location):
        return {"file": entry["file"], "relativePath": entry["relativePath"], **location, "sha256": entry["sha256"]}

    knowledge_by_model = {}
    checklists = {}
    for source in index["sources"]:
        model = source["model"]
        if model in knowledge_by_model:
            continue
        folder = (root / source["relative_path"]).parent.parent
        knowledge = {"faultCodes": [], "repairGuides": [], "notices": []}
        for prefix in ("05", "06", "09"):
            paths = sorted(p for d in folder.iterdir() if d.is_dir() and d.name.startswith(prefix)
                           for p in d.iterdir() if p.suffix.lower() in {".pptx", ".xlsx"} and not p.name.startswith("~$"))
            for path in paths:
                if prefix == "09" and path.name not in {"功能测试SOP.xlsx", "功能测试标准.xlsx", "功能测试.pptx"}:
                    continue  # Use actual functional SOPs, not equipment or separate APP manuals.
                entry = register(path, {"05": "fault-codes", "06": "repair-guides", "09": "functional-sop"}[prefix])
                if path.suffix == ".xlsx":
                    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
                    extracted_sheets = []
                    items = []
                    for ws in book:
                        rows = [{"row": n, "cells": list(row)} for n, row in enumerate(ws.iter_rows(values_only=True), 1)]
                        extracted_sheets.append({"sheet": ws.title, "rows": rows})
                        for row in rows:
                            label = row["cells"][1] if len(row["cells"]) > 1 else None
                            # Concise functional list; preserve source titles, no copied numerical test thresholds.
                            if label in {"DC输入", "AC输入", "USB口输出", "车充输出", "AC输出\nAC频率", "AC输出测试", "显示屏", "LED测试"}:
                                items.append({"code": f"CHECK_{len(items)+1:02d}", "labelEn": "", "labelZh": label,
                                              "source": reference(entry, sheet=ws.title, row=row["row"])})
                    book.close()
                    checklists[model] = items
                    extraction.append({"source": entry, "sheets": extracted_sheets})
                    continue
                slides = extract_ppt(path)
                extraction.append({"source": entry, "slides": slides})
                for slide in slides:
                    src = reference(entry, slide=slide["slide"])
                    if prefix == "05":
                        for table in slide["tables"]:
                            for row in table[1:]:
                                cells = [x["text"] for x in row["cells"]]
                                if re.fullmatch(r"F\d+", cells[0]):
                                    assert cells[0] in {f"F{i}" for i in range(10)}, (path, cells[0])
                                    knowledge["faultCodes"].append({"code": cells[0], "titleEn": "", "titleZh": cells[1],
                                        "descriptionEn": "", "descriptionZh": cells[2], "solutionEn": "", "solutionZh": "",
                                        "source": {**src, "row": row["row"]}})
                                elif any(cells[1:]):
                                    knowledge["notices"].append({"id": f"temperature-{row['row']}", "titleEn": "", "titleZh": cells[1],
                                        "descriptionEn": "", "descriptionZh": cells[2], "source": {**src, "row": row["row"]}})
                    if prefix == "06" and slide["tables"]:
                        # Keep the complete slide table together. Blank merged cells remain blank;
                        # never attach a neighbouring symptom or component to an unrelated row.
                        table_text = "\n\n".join("\n".join(" | ".join(c["text"] for c in row["cells"]) for row in table)
                                                   for table in slide["tables"])
                        title = "\n".join(slide["titles"])
                        knowledge["repairGuides"].append({"id": f"{model.lower()}-common-slide-{slide['slide']}",
                            "titleEn": "", "titleZh": title, "symptomsEn": "", "symptomsZh": "", "stepsEn": "", "stepsZh": table_text,
                            "source": src, "sourceTables": slide["tables"]})
                    if prefix == "09" and slide["tables"]:
                        items = []
                        wanted = {"AC频率", "LED照明", "显示屏"}
                        for table in slide["tables"]:
                            for row in table:
                                cells = [c["text"] for c in row["cells"]]
                                for cell_text in cells[:2]:
                                    label = cell_text.split("\n", 1)[0]
                                    if label in wanted or re.fullmatch(r"[123]、(?:DC输入|AC输入|USB口输出|车充口输出|AC交流输出)", label):
                                        items.append({"code": f"CHECK_{len(items)+1:02d}", "labelEn": "", "labelZh": label,
                                                      "source": {**src, "row": row["row"]}})
                        if items:
                            checklists[model] = items
        assert {f["code"] for f in knowledge["faultCodes"]} == {f"F{i}" for i in range(10)}, model
        assert len(knowledge["notices"]) == 2, model
        assert 3 <= len(checklists.get(model, [])) <= 8, (model, checklists.get(model))
        knowledge_by_model[model] = knowledge

    for source in index["sources"]:
        path = root / source["relative_path"]
        source = {**source, "path": str(path)}
        rows = [r for r in records if r["来源文件"] == path.name]
        verify_bom_source(source, rows)
        entry = register(path, "bom")
        base, version = source["model"], source["version"]
        key = f"{base.lower()}-{VERSION_KEYS[version]}"
        roots = [r for r in rows if str(r["BOM层级（原表）"]) == "0"]
        reasons = []
        if base == "E1000V2" and version == "原始版":
            reasons.append("filename_model_conflicts_with_source_model:E1000V2/E2000V2")
        if roots and "停用" in (roots[0]["规格型号"] or ""):
            reasons.append("root_specification_marked_停用")
        model = {"catalogKey": key, "baseModelCode": base, "bomVersion": version, "market": "EU",
                 "productSku": roots[0]["物料编码"] if roots else "", "nameEn": base, "nameZh": base,
                 "active": not reasons, "reviewReasons": reasons, "checkItems": checklists[base], "bom": [],
                 "knowledge": knowledge_by_model[base], "source": reference(entry),
                 "sourceModelCodes": list(dict.fromkeys(r["型号（原表）"] for r in rows)),
                 "rootMetadata": roots}
        if base == "E1000":
            model["reviewReasons"].append("incomplete_bom_no_product_sku_or_quantities:knowledge_and_fault_recording_only")
        catalog["models"].append(model)
        for r in rows:
            src = reference(entry, sheet=r["来源工作表"], row=r["来源行号"])
            if str(r["BOM层级（原表）"]) == "0" or r["物料编码"] in SOFTWARE:
                reason = "product_root_not_physical_part" if str(r["BOM层级（原表）"]) == "0" else "software_not_physical_part"
                catalog["excluded"].append({"reason": reason, "catalogKey": key, "sku": r["物料编码"], "source": src, "original": r})
                continue
            row_reasons = []
            qty = r["标准用量（原表）"]
            if qty is None:
                row_reasons.append("standard_quantity_missing")
            elif not isinstance(qty, (int, float)) or isinstance(qty, bool) or qty <= 0:
                row_reasons.append("standard_quantity_not_positive")
            if r["物料生命周期"] in {"停止生产", "样机", "试产"}:
                row_reasons.append("source_lifecycle:" + r["物料生命周期"])
            if r["发料方式"] == "不发料":
                row_reasons.append("physical_material_marked_不发料")
            if not model["active"]:
                row_reasons.append("parent_model_requires_review")
            bom = {"sku": r["物料编码"], "descriptionEn": "", "descriptionZh": (r["物料名称"] or "") + " | " + (r["规格型号"] or ""),
                   "unit": "pcs", "standardQty": qty, "active": not row_reasons, "reviewReasons": row_reasons,
                   "source": src, "original": r}
            model["bom"].append(bom)
            if row_reasons:
                catalog["reviewRows"].append({"catalogKey": key, "sku": bom["sku"], "reviewReasons": row_reasons, "source": src})
    counts = {"sourceBomRows": len(records), "bomSourcesVerified": len(index["sources"]), "models": len(catalog["models"]),
              "baseModels": len(knowledge_by_model), "activeModels": sum(m["active"] for m in catalog["models"]),
              "physicalCandidateRows": sum(len(m["bom"]) for m in catalog["models"]),
              "activePhysicalRows": sum(b["active"] for m in catalog["models"] for b in m["bom"]),
              "suspendedPhysicalRows": len(catalog["reviewRows"]), "excludedRows": len(catalog["excluded"]),
              "exclusionReasons": dict(Counter(x["reason"] for x in catalog["excluded"])),
              "knowledgePerBaseModel": {base: {k: len(v) for k, v in knowledge.items()} for base, knowledge in knowledge_by_model.items()}}
    assert counts["sourceBomRows"] == 414
    assert counts["exclusionReasons"] == {"product_root_not_physical_part": 8, "software_not_physical_part": 2}
    assert counts["physicalCandidateRows"] == 404
    assert counts["activePhysicalRows"] + counts["suspendedPhysicalRows"] + counts["excludedRows"] == 414
    assert len({m["catalogKey"] for m in catalog["models"]}) == 9
    for entry in catalog["sources"]:
        assert sha256(root / entry["relativePath"]) == entry["sha256"], entry["relativePath"]
    catalog["counts"] = counts
    return catalog, extraction


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--index", type=Path, default=DEFAULT_INDEX)
    parser.add_argument("--source-root", type=Path)
    args = parser.parse_args()
    catalog, extraction = build(args.index, args.source_root)
    save_json(REPO / "catalogs/jackery-eu-20261030.json", catalog)
    save_json(REPO / "artifacts/jackery-catalog/source-extraction.json", extraction)
    save_json(REPO / "artifacts/jackery-catalog/verification.json", catalog["counts"])
    print(json.dumps(catalog["counts"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
